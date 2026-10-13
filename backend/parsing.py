import csv
import io
from pathlib import Path
from .models import Quote

HEADERS={'供应商':'supplier','型号':'model','单价':'price','计价数量':'price_unit','起订量':'moq','交期':'lead_days','含税':'tax_included','税率':'tax_rate','运费':'shipping','币种':'currency'}

def parse_quotes(filename,content):
    ext=Path(filename).suffix.lower()
    if ext=='.csv':
        rows=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
        sources=[f'{filename} 第{i+2}行' for i in range(len(rows))]
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        sheet=book.active; values=list(sheet.iter_rows(values_only=True))
        if not values: raise ValueError('工作表为空')
        rows=[dict(zip(values[0],r)) for r in values[1:] if any(x is not None for x in r)]
        sources=[f'{filename} / {sheet.title} 第{i+2}行' for i in range(len(rows))]
    else: raise ValueError('报价导入支持 CSV 和 XLSX；PDF/TXT 可上传为规格知识资料。')
    if not 1<=len(rows)<=20: raise ValueError('每次导入 1–20 条报价')
    quotes=[]
    for i,(raw,source) in enumerate(zip(rows,sources)):
        clean={HEADERS.get(str(k).strip(),str(k).strip()):v for k,v in raw.items() if k is not None and v not in ('',None)}
        clean={k:v for k,v in clean.items() if k in Quote.model_fields}
        if 'tax_included' in clean:
            v=str(clean['tax_included']).strip().lower()
            if v not in ['true','false','1','0','是','否']: raise ValueError('含税字段仅接受 true/false、1/0 或 是/否')
            clean['tax_included']=v in ['true','1','是']
        clean['id']=f'q{i+1}';clean['source']=source
        quotes.append(Quote.model_validate(clean).model_dump(mode='json'))
    return quotes

def parse_knowledge(filename,content,models):
    ext=Path(filename).suffix.lower()
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(content))
        if len(reader.pages)>80: raise ValueError('知识文件最多 80 页')
        pages=[p.extract_text() or '' for p in reader.pages]
        if not any(x.strip() for x in pages): raise ValueError('扫描 PDF 没有文本层；第一版不支持 OCR，请上传文本型 PDF。')
    elif ext in ['.txt','.md']: pages=[content.decode('utf-8-sig')]
    else: raise ValueError('知识资料支持文本型 PDF、TXT、Markdown')
    import uuid
    docs=[]
    for page,text in enumerate(pages,1):
        for i,start in enumerate(range(0,len(text),1000)):
            docs.append({'id':uuid.uuid4().hex,'title':Path(filename).name,'page':page,'chunk':i,'models':models,'text':text[start:start+1200],'specs':{},'demo':False})
    if len(docs)>200: raise ValueError('文件分块超过 200 个，请缩小资料范围')
    return docs
